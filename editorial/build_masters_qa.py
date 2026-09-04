#!/usr/bin/env python3
"""Build the Richard Masters Chapter 01 Q&A draft in TFOR gold-standard format.

Spec: reference_tfor_docx_format.md
- Letter page, 1in margins
- Normal style: Georgia 11pt, 1.15 line, 10pt after, no kerning/ligatures
- Title bold 14pt; metadata italic 10pt #7A756D; section headings bold 11pt
- Questions bold with 0.3in hanging indent; answers plain
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

STONE = "7A756D"

OUT = ("/Users/jack/firstownersreference/editorial/contributor-docs/"
       "Richard Masters - Chapter 01 Q&A (draft for approval).docx")


def setup(doc):
    for s in doc.sections:
        s.page_width = Inches(8.5)
        s.page_height = Inches(11)
        for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(s, m, Inches(1))

    normal = doc.styles["Normal"]
    normal.font.name = "Georgia"
    normal.font.size = Pt(11)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), "Georgia")
    for tag in ("w:kern", "w14:ligatures"):
        for el in rpr.findall(qn(tag)):
            rpr.remove(el)
    ppr = normal.element.get_or_add_pPr()
    spacing = ppr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        ppr.append(spacing)
    spacing.set(qn("w:after"), "200")
    spacing.set(qn("w:line"), "276")
    spacing.set(qn("w:lineRule"), "auto")


def para(doc, text="", bold=False, italic=False, size=None, colour=None,
         hanging=False, space_before=None):
    p = doc.add_paragraph()
    if space_before is not None:
        p.paragraph_format.space_before = Pt(space_before)
    if hanging:
        ppr = p._p.get_or_add_pPr()
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "432")
        ind.set(qn("w:hanging"), "432")
        ppr.append(ind)
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        if size:
            r.font.size = Pt(size)
        if colour:
            r.font.color.rgb = RGBColor.from_string(colour)
    return p


TITLE = "Richard Masters, Director - Master Yachts"
META = ("Chapter 01, The reality of ownership. Guest Q&A, revised draft for approval, "
        "4 September 2026.")

INTRO = (
    "Master Yachts has managed yachts in operation for 20 years, more than 50 of them "
    "between 20 and 148 metres, and represents owners through new build on a set fee "
    "agreed monthly or quarterly. It provides services discussed in this chapter, so the "
    "independence test set out in the essay above applies to it as to any firm. We put six "
    "questions to Richard Masters on the percentage-of-value running cost model, the "
    "year-one cost surprises that recur, where the crew budget moves, use intensity among "
    "first-time owners, the independence test applied to his own firm, and the first three "
    "years as he would put them to a first-time buyer. His answers are published as given, "
    "lightly edited for length and clarity."
)

# (question, [answer paragraphs], pull quote or None)
QA = [
    (
        "Cost reality across the Master Yachts fleet: what percentage of purchase price "
        "actually holds for first-time owners in years one to three, and where does it "
        "land by yacht size and use intensity?",
        [
            "We have rarely agreed with a maintenance and running cost model based on a "
            "percentage of the vessel’s value. It is a fallacy. A way to sell. A broker "
            "creation. It is not operationally logical at all.",
            "The considerations are the number of crew, the operational requirements, the "
            "quality of construction, the weeks of use, the preventative maintenance "
            "requirement, the owner’s standard of presentation, consumable burn rates, and "
            "a good many other contributing factors.",
        ],
        "A running cost model based on a percentage of the vessel’s value is a fallacy. "
        "A way to sell. A broker creation.",
    ),
    (
        "The single biggest year-one cost surprise that recurs across the owners you have "
        "managed for. Where does the headline budget most often break down, and what is "
        "the typical magnitude of the gap?",
        [
            "Crew turnover until the boat is settled. Insurance costs, particularly medical. "
            "Fuel costs, which move with global fluctuations. And warranty travel costs, "
            "which nobody considers at the time and which can run to hundreds of thousands.",
            "Warranty travel is not selective by yacht size or by type of owner. Imagine a "
            "beautifully constructed Dutch yacht relocated to New Zealand, a German quality "
            "build operating in Spain, or a superb Turkish build sitting in a French "
            "shipyard. The figures are staggering and are very often forgotten.",
        ],
        "Warranty travel costs, which nobody considers at the time, can run to hundreds "
        "of thousands.",
    ),
    (
        "Crew is documented at 30 to 40 percent of operating cost. From inside fleet "
        "management, where does that cost actually go that the owner does not see on the "
        "budget summary?",
        [
            "In our experience it is always minutely detailed on the initial budget. It is "
            "the potential turnover that can be the large cost surprise.",
            "Turnover can be as much as 50 percent of a crew of 5 to 50 in the first year "
            "of operation. It depends very much on leadership and on the owner’s taste. It "
            "is not a hard fact, but it is a cost surprise that is often left out of a "
            "pre-launch operational budget.",
        ],
        None,
    ),
    (
        "The widely cited pattern that many first-time owners use their yacht only four to "
        "eight weeks a year. Does it hold across your fleet, and what do owners with 12 "
        "or more weeks of use a year do differently, in operational and financial terms?",
        [
            "First-time owners are generally younger and have very busy work schedules. "
            "These younger owners are also extremely savvy. They ask questions and they "
            "demand answers.",
            "Owners able to spend 12 or more weeks aboard are generally experienced and "
            "very aware of the costs of use.",
        ],
        None,
    ),
    (
        "The independence test in this chapter, applied to your own firm. Master Yachts "
        "represents owners through new build and manages yachts in operation, both on "
        "recurring fees. How do the two sit together, and how is the fee basis for each "
        "set and disclosed at engagement?",
        [
            "We are new build representatives first and operational managers second. Most "
            "of what I have said here comes from the operational side, from looking after "
            "yachts once they are in service, and I would not want it read as anything "
            "more than that.",
            "New build representation costs depend on complexity, the size of the team on "
            "site, the duration, and similar factors. It is always a set fee, agreed monthly "
            "or quarterly. Operational management is governed much more by industry norms, "
            "set against a proper risk analysis of the intended use.",
        ],
        None,
    ),
    (
        "The honest case for ownership versus charter, told from someone who has seen both "
        "close to. For a first-time UHNW buyer reading this, what is the one piece of "
        "operational reality about the first three years of ownership that you would put in "
        "front of them?",
        [
            "Some owners simply do not have the time in their schedules to use the vessel as "
            "much as they would like, and they wish to offset operational costs where they "
            "can. We have yachts that are absolutely not available for charter at any price. "
            "Some are available selectively, but not in high season. Others where the owners "
            "will work their schedules around high-season bookings.",
            "On the first three years: years one and two are about very sound warranty "
            "management, and ensuring excellent engagement with the builder and the "
            "suppliers. Year three is reality. By then the style of use is known, the costs "
            "can be well controlled and understood, and there are few surprises.",
        ],
        "Years one and two are about very sound warranty management. Year three is reality.",
    ),
]

QUERIES = [
    "The introduction and question five have been rewritten from your note of 4 September. "
    "The firm is now described as operational manager and new build representative, and "
    "the answer to question five notes that the answers come from the operating side. "
    "Please check both read as you want them.",

    "Warranty travel and crew turnover run in your words, with your caveats attached. The "
    "byline reads “Richard Masters, Director - Master Yachts”. Please tell us if you would "
    "like either otherwise.",
]

FOOTER = (
    "Draft for review. Nothing is published until you have approved this text. The edits "
    "are light, for length and clarity only, and the substance is yours. With thanks, Jack "
    "MacNally, on behalf of The First Owner’s Reference, September 2026."
)


def main():
    doc = Document()
    setup(doc)

    para(doc, TITLE, bold=True, size=14)
    para(doc, META, italic=True, size=10, colour=STONE)

    para(doc, "Introduction as it will appear", bold=True, space_before=6)
    para(doc, INTRO)

    para(doc, "Questions and answers", bold=True, space_before=6)
    for i, (q, answers, pull) in enumerate(QA, start=1):
        para(doc, f"{i}. {q}", bold=True, hanging=True)
        for a in answers:
            para(doc, a)
        if pull:
            para(doc, "Pull quote: " + pull, italic=True, colour=STONE)

    para(doc, "Queries before publication", bold=True, space_before=6)
    for q in QUERIES:
        para(doc, q)

    para(doc, FOOTER, italic=True, size=10, colour=STONE, space_before=6)

    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
