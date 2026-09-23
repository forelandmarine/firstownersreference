#!/usr/bin/env python3
"""Build the Benjamin Maltby Chapter 05 question set in TFOR contributor-doc format.

Matches the layout of the existing question docs (Hugo Morgan-Harris - Chapter 04.docx):
Letter page, 1in margins, Georgia 11pt Normal, bold 14pt title, italic 10pt stone metadata,
numbered questions with bold number and 0.3in hanging indent, italic stone closing note.
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

STONE = "7A756D"
OUT = ("/Users/jack/firstownersreference/editorial/contributor-docs/"
       "Benjamin Maltby - Chapter 05.docx")

TITLE = "Benjamin Maltby, Keystone Law"
META = (
    "Chapter 05, New build versus brokerage. Partner, Keystone Law; barrister, called in 2000, "
    "formerly at Ince & Co and MatrixLloyd; advises on superyacht construction, purchase and sale, "
    "finance, and operational matters including tax, insurance and employment. Instructed on a fee "
    "basis by the party he acts for, which across the practice has included owners, finance "
    "providers, yards and suppliers; not paid contingent on any transaction closing."
)

QUESTIONS = [
    "The European yards that matter do not contract on SAJ or NEWBUILDCON. They contract on "
    "in-house templates drafted by their own counsel, with English law and London arbitration "
    "attached. When a first-time buyer’s team receives that first draft, where does the balance "
    "of risk sit, and which clauses do you open first?",

    "Refund guarantees are meant to secure the stage payments a buyer pays forward. Nobiskrug in "
    "2024 and Italian Sea Group in 2026 have shown what happens when a yard’s financial condition "
    "changes mid-build. In practice, how much protection does a refund guarantee give, what "
    "separates a reliable one from a weak one, and what should be agreed at heads of terms about "
    "a distressed-yard scenario before it arises?",

    "Title during construction. Some yards retain title until delivery; others pass it "
    "progressively as stages are paid. For an owner who has paid 60 percent of the contract price "
    "forward, what does each position mean if the yard fails, and what security would you want in "
    "place alongside it?",

    "The contracting party. Winch Design v Le Souef (2025) turned on whether the individual or "
    "his special purpose vehicle was the party to the contract. What does a buyer need in place, "
    "and by when, for an ownership structure to hold in a dispute rather than only on paper?",

    "Change-order procedure, the liquidated damages cap and warranty length are the points where "
    "the chapter says inexperience gets priced at signature. Which of these do owners most "
    "consistently give away, and what does a well-negotiated position on it look like?",

    "On the brokerage side, the MYBA MOA is signed at heads of terms with a 10 percent deposit, "
    "often under time pressure. What do first-time buyers most often misunderstand about what the "
    "MOA commits them to, and at what point should their lawyer be in the room?",

    "The chapter treats VAT regime, flag and ownership structure as one decision with three "
    "parts, to be settled before contracts are exchanged. From your practice, what happens when "
    "that decision is left until after closing, and what does unpicking it cost?",
]

CLOSE = (
    "Seven suggested questions, offered as a starting point. Please reshape any of them, and add "
    "an angle we have not asked. With thanks, Jack MacNally, on behalf of The First Owner’s "
    "Reference, September 2026."
)


def setup(doc):
    for s in doc.sections:
        s.page_width = Inches(8.5)
        s.page_height = Inches(11)
        for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(s, m, Inches(1))
    normal = doc.styles["Normal"]
    normal.font.name = "Georgia"
    normal.font.size = Pt(11)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), "Georgia")
    for tag in ("w:kern", "w14:ligatures"):
        for el in rpr.findall(qn(tag)):
            rpr.remove(el)
    ppr = normal.element.get_or_add_pPr()
    spacing = ppr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        ppr.append(spacing)
    spacing.set(qn("w:after"), "200")
    spacing.set(qn("w:line"), "276")
    spacing.set(qn("w:lineRule"), "auto")


def para(doc, after_pt, hanging=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after_pt)
    if hanging:
        ppr = p._p.get_or_add_pPr()
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "432")
        ind.set(qn("w:hanging"), "432")
        ppr.append(ind)
    return p


def main():
    doc = Document()
    setup(doc)

    p = para(doc, 6)
    r = p.add_run(TITLE)
    r.bold = True
    r.font.size = Pt(14)
    r.font.name = "Georgia"

    p = para(doc, 12)
    r = p.add_run(META)
    r.italic = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(STONE)

    for i, q in enumerate(QUESTIONS, 1):
        p = para(doc, 10, hanging=True)
        n = p.add_run(f"{i}.  ")
        n.bold = True
        p.add_run(q)

    p = para(doc, 10)
    r = p.add_run(CLOSE)
    r.italic = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(STONE)

    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
